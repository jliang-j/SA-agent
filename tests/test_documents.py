import io

import pytest
from fastapi.testclient import TestClient
from docx import Document as WordDocument
from langchain_core.embeddings import Embeddings
from pypdf import PdfWriter
from pypdf.generic import DecodedStreamObject, DictionaryObject, NameObject

import sa_agent.main as main
from sa_agent import documents
from sa_agent.retrieval.ingest import ingest_document_file
from sa_agent.retrieval.loader import load_and_split_document
from sa_agent.retrieval.store import create_persistent_vector_store


class LocalEmbeddings(Embeddings):
    """使用本地固定维度向量验证真实Chroma写入，不调用外部服务。"""

    def embed_documents(self, texts):
        """生成用于测试的文档向量。"""
        return [[1.0, float(len(text)), 0.5] for text in texts]

    def embed_query(self, text):
        """生成测试查询向量。"""
        return self.embed_documents([text])[0]


def pdf_bytes(texts):
    """生成仅供解析测试的内存PDF，各页可有文本或为空。"""
    writer = PdfWriter()
    font = DictionaryObject({NameObject("/Type"): NameObject("/Font"),
                             NameObject("/Subtype"): NameObject("/Type1"),
                             NameObject("/BaseFont"): NameObject("/Helvetica")})
    for text in texts:
        page = writer.add_blank_page(width=400, height=400)
        if text:
            page[NameObject("/Resources")] = DictionaryObject({
                NameObject("/Font"): DictionaryObject({NameObject("/F1"): font})})
            stream = DecodedStreamObject()
            stream.set_data(f"BT /F1 12 Tf 20 300 Td ({text}) Tj ET".encode("ascii"))
            page[NameObject("/Contents")] = stream
    buffer = io.BytesIO()
    writer.write(buffer)
    return buffer.getvalue()


def docx_bytes(empty=False):
    """构造含正文和普通表格的内存DOCX，或无正文的有效文件。"""
    document = WordDocument()
    if not empty:
        document.add_paragraph("企业应保留研究证据。")
        table = document.add_table(rows=1, cols=2)
        table.cell(0, 0).text = "复核频率"
        table.cell(0, 1).text = "每年一次"
    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()


@pytest.fixture
def client(tmp_path, monkeypatch):
    """隔离文档和向量库目录，保留真实解析与入库，仅替换Embedding。"""
    monkeypatch.setattr(main, "DOCUMENTS_DIR", tmp_path / "documents")
    monkeypatch.setattr(main, "VECTOR_STORE_DIR", tmp_path / "chroma")
    real_ingest = documents.ingest_document_file
    real_store = documents.create_persistent_vector_store

    def ingest(*args, **kwargs):
        """用本地向量执行真实导入。"""
        return real_ingest(*args, **kwargs, embeddings=LocalEmbeddings())

    def store(**kwargs):
        """打开同一测试库以读取实际片段统计。"""
        return real_store(**kwargs, embeddings=LocalEmbeddings())

    monkeypatch.setattr(documents, "ingest_document_file", ingest)
    monkeypatch.setattr(documents, "create_persistent_vector_store", store)
    with TestClient(main.app) as value:
        yield value


@pytest.mark.parametrize("name", ["report.txt", "report.md"])
def test_upload_text(client, name):
    """上传UTF-8文本并校验实际入库与文档列表。"""
    response = client.post("/documents/upload", files={"file": (name, "企业应加强数据治理。".encode())})
    assert response.status_code == 201, response.text
    assert response.json()["status"] == "imported"
    listing = client.get("/documents").json()["documents"]
    assert listing[0]["filename"] == name
    assert listing[0]["status"] == "imported"
    assert listing[0]["chunks_count"] == 1
    assert "source_path" not in listing[0]


def test_pdf_pages(client):
    """文本PDF保留原始页码，片段编号跨页唯一。"""
    response = client.post("/documents/upload", files={"file": ("report.pdf", pdf_bytes(["First page", "Second page"]))})
    assert response.status_code == 201, response.text
    parsed = load_and_split_document(main.DOCUMENTS_DIR / "report.pdf")
    assert [doc.metadata["page"] for doc in parsed] == [1, 2]
    assert [doc.metadata["chunk_index"] for doc in parsed] == [0, 1]
    assert client.get("/documents").json()["documents"][0]["chunks_count"] == 2


@pytest.mark.parametrize("name,content,code", [
    ("../bad.txt", b"data", 400), ("bad.doc", b"data", 400),
    ("bad.docx", b"not a docx", 422),
    ("NUL.txt", b"data", 400), ("bad.txt", b"", 422),
    ("bad.txt", b"\xff\xfe\xff", 422), ("bad.pdf", b"not a pdf", 422),
])
def test_bad_upload(client, name, content, code):
    """拒绝不支持的格式、路径和无效文件内容。"""
    response = client.post("/documents/upload", files={"file": (name, content)})
    assert response.status_code == code, response.text


@pytest.mark.parametrize("name", ["report.docx", "report.DOCX"])
def test_upload_docx(client, name):
    """DOCX上传后正文和表格实际入库，列表显示格式及片段数量。"""
    response = client.post("/documents/upload", files={"file": (name, docx_bytes())})
    assert response.status_code == 201, response.text
    assert response.json()["status"] == "imported"
    listing = client.get("/documents").json()["documents"]
    assert listing[0]["format"] == "DOCX"
    assert listing[0]["chunks_count"] == 1
    stored = create_persistent_vector_store(
        persist_directory=main.VECTOR_STORE_DIR, collection_name=main.COLLECTION_NAME,
        embeddings=LocalEmbeddings(),
    ).get()
    assert "企业应保留研究证据" in stored["documents"][0]
    assert "复核频率\t每年一次" in stored["documents"][0]
    assert stored["metadatas"][0]["source"] == name
    assert "page" not in stored["metadatas"][0]


def test_empty_docx_upload(client):
    """无正文DOCX记录失败，重试仍返回明确错误而非零片段成功。"""
    response = client.post("/documents/upload", files={"file": ("empty.docx", docx_bytes(empty=True))})
    assert response.status_code == 422
    assert client.get("/documents").json()["documents"][0]["status"] == "failed"
    assert client.post("/documents/retry", json={"filename": "empty.docx"}).status_code == 422


def test_docx_retry(client, monkeypatch):
    """DOCX向量服务失败后可重试，重复导入不增加片段。"""
    original = documents.ingest_document_file

    def fail(*args, **kwargs):
        """模拟暂时不可用的向量服务。"""
        raise RuntimeError("temporary failure")

    monkeypatch.setattr(documents, "ingest_document_file", fail)
    assert client.post("/documents/upload", files={"file": ("retry.docx", docx_bytes())}).status_code == 500
    monkeypatch.setattr(documents, "ingest_document_file", original)
    for _ in range(2):
        assert client.post("/documents/retry", json={"filename": "retry.docx"}).status_code == 200
    assert client.get("/documents").json()["documents"][0]["chunks_count"] == 1


def test_no_pdf_text(client):
    """无可提取正文时明确失败，不写入零片段成功状态。"""
    response = client.post("/documents/upload", files={"file": ("blank.pdf", pdf_bytes([""]))})
    assert response.status_code == 422
    assert "OCR" in response.json()["detail"]
    assert client.get("/documents").json()["documents"][0]["status"] == "failed"


def test_pdf_page_limit(client):
    """超过页数上限的PDF在向量化前被拒绝。"""
    response = client.post("/documents/upload", files={"file": ("long.pdf", pdf_bytes([""] * 201))})
    assert response.status_code == 422
    assert "200" in response.json()["detail"]


def test_encrypted_pdf(client):
    """加密PDF返回明确错误，不尝试解密。"""
    writer = PdfWriter()
    writer.add_blank_page(width=400, height=400)
    writer.encrypt("test-only-password")
    buffer = io.BytesIO()
    writer.write(buffer)
    response = client.post("/documents/upload", files={"file": ("locked.pdf", buffer.getvalue())})
    assert response.status_code == 422
    assert "加密" in response.json()["detail"]


def test_duplicate_upload(client):
    """同名上传返回冲突，原文件及片段保持不变。"""
    assert client.post("/documents/upload", files={"file": ("same.txt", b"original")}).status_code == 201
    response = client.post("/documents/upload", files={"file": ("same.txt", b"replacement")})
    assert response.status_code == 409
    assert (main.DOCUMENTS_DIR / "same.txt").read_bytes() == b"original"
    assert client.get("/documents").json()["documents"][0]["chunks_count"] == 1


def test_upload_size(client, monkeypatch):
    """超出大小限制时不保存文件。"""
    monkeypatch.setattr(documents, "MAX_BYTES", 4)
    response = client.post("/documents/upload", files={"file": ("big.txt", b"12345")})
    assert response.status_code == 413
    assert not (main.DOCUMENTS_DIR / "big.txt").exists()


def test_retry(client, monkeypatch):
    """导入服务失败后保留文件和失败状态，可重试成功且不重复入库。"""
    original = documents.ingest_document_file

    def fail(*args, **kwargs):
        """模拟向量服务暂时失败。"""
        raise RuntimeError("temporary failure")

    monkeypatch.setattr(documents, "ingest_document_file", fail)
    response = client.post("/documents/upload", files={"file": ("retry.txt", b"content")})
    assert response.status_code == 500
    assert (main.DOCUMENTS_DIR / "retry.txt").exists()
    assert client.get("/documents").json()["documents"][0]["status"] == "failed"
    monkeypatch.setattr(documents, "ingest_document_file", original)
    assert client.post("/documents/retry", json={"filename": "retry.txt"}).status_code == 200
    assert client.post("/documents/retry", json={"filename": "retry.txt"}).status_code == 200
    assert client.get("/documents").json()["documents"][0]["chunks_count"] == 1


def test_existing_files(client):
    """原有文件没有状态记录时仍可列出，并支持从列表导入。"""
    main.DOCUMENTS_DIR.mkdir()
    (main.DOCUMENTS_DIR / "existing.md").write_text("已有资料", encoding="utf-8")
    assert client.get("/documents").json()["documents"][0]["status"] == "not_imported"
    assert client.post("/documents/retry", json={"filename": "existing.md"}).status_code == 200
    assert client.get("/documents").json()["documents"][0]["status"] == "imported"


def test_missing_retry(client):
    """重试不存在的文档时返回404。"""
    assert client.post("/documents/retry", json={"filename": "missing.txt"}).status_code == 404


def test_failed_embedding_preserves_chunks(tmp_path):
    """再次向量化失败时，不提前删除已有片段。"""
    path = tmp_path / "keep.txt"
    path.write_text("old content", encoding="utf-8")
    kwargs = {"persist_directory": tmp_path / "chroma", "collection_name": "preserve_test"}
    ingest_document_file(path, **kwargs, embeddings=LocalEmbeddings())

    class BrokenEmbeddings(LocalEmbeddings):
        """模拟文档向量化失败。"""

        def embed_documents(self, texts):
            """在写入前抛出异常。"""
            raise RuntimeError("embedding failed")

    path.write_text("new content", encoding="utf-8")
    with pytest.raises(RuntimeError):
        ingest_document_file(path, **kwargs, embeddings=BrokenEmbeddings())
    stored = create_persistent_vector_store(**kwargs, embeddings=LocalEmbeddings()).get()
    assert stored["documents"] == ["old content"]
