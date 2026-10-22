import pytest
from docx import Document as WordDocument

import sa_agent.retrieval.ingest as ingest
from sa_agent.retrieval.loader import load_and_split_document, load_docx_text


def test_docx_order(tmp_path):
    """正文与普通表格保持原始顺序，表格保留行列分隔。"""
    path = tmp_path / "report.docx"
    document = WordDocument()
    document.add_heading("治理报告", level=1)
    document.add_paragraph("表格前的说明。")
    table = document.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "措施"
    table.cell(0, 1).text = "频率"
    table.cell(1, 0).text = "来源复核"
    table.cell(1, 1).text = "每年一次"
    document.add_paragraph("表格后的说明。")
    document.save(path)

    text = load_docx_text(path)
    assert text.index("表格前") < text.index("措施\t频率") < text.index("表格后")
    assert "来源复核\t每年一次" in text
    assert text.startswith("治理报告")


@pytest.mark.parametrize("suffix", [".docx", ".DOCX"])
def test_docx_chunks(tmp_path, suffix):
    """DOCX复用现有切分，编号唯一连续，保留文件名但不伪造页码。"""
    path = tmp_path / f"report{suffix}"
    document = WordDocument()
    document.add_paragraph("企业需要保留来源与原文证据。" * 100)
    document.save(path)

    chunks = load_and_split_document(path)
    assert len(chunks) > 1
    assert [item.metadata["chunk_index"] for item in chunks] == list(range(len(chunks)))
    assert all(item.metadata["source"] == path.name for item in chunks)
    assert all("page" not in item.metadata for item in chunks)


def test_empty_docx(tmp_path, monkeypatch):
    """空DOCX不进入向量化，不产生零片段导入成功。"""
    path = tmp_path / "empty.docx"
    WordDocument().save(path)

    def unexpected(**kwargs):
        """确保空文档不会创建向量库或调用模型。"""
        pytest.fail("不应创建向量库")

    monkeypatch.setattr(ingest, "create_persistent_vector_store", unexpected)
    assert load_and_split_document(path) == []
    with pytest.raises(ValueError, match="有效正文"):
        ingest.ingest_document_file(path)


def test_invalid_docx(tmp_path):
    """伪造扩展名或损坏的DOCX返回可理解的错误。"""
    path = tmp_path / "broken.docx"
    path.write_text("这不是Word文件", encoding="utf-8")
    with pytest.raises(ValueError, match="DOCX解析失败"):
        load_and_split_document(path)


def test_missing_docx(tmp_path):
    """不存在的文件返回文件不存在错误。"""
    with pytest.raises(FileNotFoundError):
        load_docx_text(tmp_path / "missing.docx")


def test_old_doc(tmp_path):
    """明确拒绝旧式DOC格式，不尝试当作DOCX解析。"""
    path = tmp_path / "old.doc"
    path.write_bytes(b"unsupported")
    with pytest.raises(ValueError, match="只支持"):
        load_and_split_document(path)
